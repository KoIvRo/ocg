from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from PIL import Image, ImageDraw


def create_images(output_dir):
    signature = Image.new("RGBA", (500, 160), "white")
    draw = ImageDraw.Draw(signature)
    draw.line([(20, 120), (90, 45), (135, 125), (205, 35), (250, 120), (470, 85)], fill="navy", width=7)
    signature.save(output_dir / "signature.png")

    stamp = Image.new("RGBA", (300, 300), "white")
    draw = ImageDraw.Draw(stamp)
    draw.ellipse((20, 20, 280, 280), outline="darkred", width=9)
    draw.ellipse((55, 55, 245, 245), outline="darkred", width=5)
    draw.text((105, 130), "WEDDING", fill="darkred")
    stamp.save(output_dir / "stamp.png")


def main():
    output_dir = Path("checks/2")
    output_dir.mkdir(parents=True, exist_ok=True)
    create_images(output_dir)
    document = Document()

    title = document.add_heading("Приглашение на свадьбу", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("Дорогой гость", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Мы будем рады видеть Вас ").bold = True
    paragraph.add_run("на торжестве, посвященном нашей свадьбе.").italic = True

    document.add_heading("Сведения о приглашении", level=2)
    table = document.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    fields = [
        ("ФИО", "{{ full_name }}"),
        ("Адрес", "{{ address }}"),
        ("Дата", "{{ date }}"),
        ("Время", "{{ time }}"),
        ("Место", "{{ place }}"),
    ]
    for row, (label, value) in zip(table.rows, fields):
        row.cells[0].text = label
        row.cells[1].text = value

    document.add_heading("Просим", level=2)
    for text in ("подтвердить участие", "сообщить о пищевых предпочтениях"):
        document.add_paragraph(text, style="List Bullet")

    document.add_heading("Порядок мероприятия", level=2)
    for text in ("Сбор гостей", "Церемония", "Праздничный ужин"):
        document.add_paragraph(text, style="List Number")

    document.add_paragraph("Подпись")
    document.add_picture(str(output_dir / "signature.png"), width=Inches(2.2))
    document.add_paragraph("Печать")
    document.add_picture(str(output_dir / "stamp.png"), width=Inches(1.4))

    document.save(output_dir / "invitation_template.docx")


if __name__ == "__main__":
    main()
